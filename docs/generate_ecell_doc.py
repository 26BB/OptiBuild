import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_callout(doc, text, title="KEY STRATEGIC TAKEAWAY"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "EFF6FF")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border: thick blue; top/bottom/right: none
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="2563EB"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"📌 {title}\n")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = RGBColor(30, 58, 138) # Navy
    
    run_body = p.add_run(text)
    run_body.font.name = "Calibri"
    run_body.font.size = Pt(10)
    run_body.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_table(table, header_bg="1E3A8A", alt_row_bg="F8FAFC", col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        # Prevent row split across pages
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        is_header = (i == 0)
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            
        for j, cell in enumerate(row.cells):
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            if col_widths and j < len(col_widths):
                cell.width = col_widths[j]
                
            if is_header:
                set_cell_background(cell, header_bg)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(4)
                    for r in p.runs:
                        r.font.bold = True
                        r.font.name = "Calibri"
                        r.font.size = Pt(9.5)
                        r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                bg = alt_row_bg if (i % 2 == 1) else "FFFFFF"
                set_cell_background(cell, bg)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    for r in p.runs:
                        r.font.name = "Calibri"
                        r.font.size = Pt(9)
                        r.font.color.rgb = RGBColor(51, 65, 85)

def build_document(output_path):
    doc = docx.Document()
    
    # Page setup - 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    # Document Header & Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_pre = title_p.add_run("THE COLLEGE INCUBATION & E-CELL STRATEGIC PLAYBOOK\n")
    run_pre.font.name = "Calibri"
    run_pre.font.size = Pt(11)
    run_pre.font.bold = True
    run_pre.font.color.rgb = RGBColor(37, 99, 235) # Blue accent
    
    run_title = title_p.add_run("Master Guide to Software, Cloud & AI Credits for Student Startups")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42) # Slate dark
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(2)
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Institutional Framework, Direct Application Routes, AI Infrastructure & Operational Software Grants\nAuthor / Role Context: President, Entrepreneurship Cell (E-Cell) / Campus Incubation Center")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)
    
    add_callout(
        doc,
        "As President of the E-Cell, your core objective is to eliminate runway friction for student-led startups. "
        "Global technology leaders (Amazon, Microsoft, Google, OpenAI, Freshworks, Zoho) allocate millions of dollars "
        "in zero-equity promotional credits to student founders. By registering your college incubation cell as an "
        "official Institutional Provider, you can grant every campus venture up to $150,000+ in software runway at zero financial cost to the university.",
        "EXECUTIVE SUMMARY FOR E-CELL PRESIDENTS"
    )
    
    # -------------------------------------------------------------
    # SECTION 1
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r = h1.add_run("1. Cloud Compute, Hosting & Infrastructure Credits")
    r.font.name = "Calibri"
    r.font.color.rgb = RGBColor(30, 58, 138)
    
    p = doc.add_paragraph(
        "Modern student applications require production cloud compute, low-latency API hosting, managed databases, "
        "and continuous deployment pipelines. The table below outlines the major infrastructure programs available."
    )
    p.paragraph_format.space_after = Pt(8)
    
    t1 = doc.add_table(rows=6, cols=5)
    headers = ["Program Name", "Credit Value", "Eligible Services", "Who Applies", "Key Value for Campus Startups"]
    for j, h in enumerate(headers):
        t1.cell(0, j).text = h
        
    infra_data = [
        ("AWS Activate\n(Founders vs. Portfolio)", "$1,000 to $25,000 (Tiered)", "EC2, S3, RDS, Lambda, SageMaker, Bedrock", "Founders directly OR via Incubation Cell Org ID", "Portfolio tier grants $10,000-$25,000 if your college E-Cell is an AWS Activate Provider."),
        ("Microsoft for Startups\nFounders Hub", "$1,000 up to $150,000", "Azure Compute, Azure OpenAI, CosmosDB, GitHub", "Student founders directly (No incubator needed for Tier 1)", "Fastest approval (3-5 days). Bundles GitHub Enterprise + $2,500 in direct OpenAI tokens."),
        ("Google for Startups\nCloud Program", "$2,000 (Start) up to $100,000+ (Scale)", "Compute Engine, Cloud Run, BigQuery, Vertex AI", "Founders or Incubation Cell Partner", "High-tier credits for AI-first startups (up to $350k for Vertex AI & Gemini APIs)."),
        ("GitHub for Startups", "20 Enterprise Seats (1 yr free)", "GitHub Enterprise, Copilot, CI/CD Actions Runners", "Incubated startups via College Partner", "Essential for collaborative codebases, automated testing pipelines, and branch protection."),
        ("MongoDB for Startups", "Up to $5,000 in Atlas Credits", "MongoDB Atlas Database, Vector Search, Technical 1:1", "Direct founder application", "Standard document database for web apps, rapid JSON prototyping, and user data.")
    ]
    
    for i, row_data in enumerate(infra_data):
        for j, text in enumerate(row_data):
            t1.cell(i+1, j).text = text
            
    format_table(t1, col_widths=[Inches(1.4), Inches(1.2), Inches(1.3), Inches(1.3), Inches(1.5)])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # -------------------------------------------------------------
    # SECTION 2
    # -------------------------------------------------------------
    h2 = doc.add_heading(level=1)
    r2 = h2.add_run("2. Artificial Intelligence, LLMs & Foundation Model Credits")
    r2.font.name = "Calibri"
    r2.font.color.rgb = RGBColor(30, 58, 138)
    
    p = doc.add_paragraph(
        "AI is the backbone of next-generation startups. Whether performing document intelligence, multi-modal site analysis, "
        "or automated agent workflows, token consumption can quickly drain student budgets. These programs offset model costs:"
    )
    p.paragraph_format.space_after = Pt(8)
    
    t2 = doc.add_table(rows=5, cols=4)
    h2_cols = ["AI Ecosystem", "Promotional Grants", "Accessible Models & Tools", "Application Strategy"]
    for j, h in enumerate(h2_cols):
        t2.cell(0, j).text = h
        
    ai_data = [
        ("OpenAI for Startups\n(via Microsoft)", "$2,500 direct credits (scales to $25k)", "GPT-4o, GPT-4o-mini, Whisper, Embeddings", "Available immediately upon Level 2 verification in Microsoft for Startups Founders Hub."),
        ("Anthropic Claude\n(via AWS Bedrock)", "Funded through AWS Activate credits", "Claude 3.5 Sonnet, Claude 3.5 Haiku, Claude 3 Opus", "Request model access inside AWS Console; usage consumes AWS Activate promotional credits."),
        ("Google Cloud Vertex AI\n& Gemini Track", "Up to $350,000 in Vertex AI credits", "Gemini 1.5 Pro, Gemini Flash, Imagen 3, Codey", "Apply for Google for Startups 'AI-First' cohort with working prototype and architecture diagram."),
        ("Groq / Together AI /\nDeepInfra Fast Inference", "$500 to $2,000 developer credits", "Llama 3, DeepSeek, Qwen 2.5, Mixtral", "Apply via developer portal. Provides sub-second inference speed for internal automation and agents.")
    ]
    for i, row_data in enumerate(ai_data):
        for j, text in enumerate(row_data):
            t2.cell(i+1, j).text = text
            
    format_table(t2, col_widths=[Inches(1.5), Inches(1.4), Inches(1.8), Inches(2.0)])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # -------------------------------------------------------------
    # SECTION 3
    # -------------------------------------------------------------
    h3 = doc.add_heading(level=1)
    r3 = h3.add_run("3. Customer Support, Live Helpdesk & Sales CRM")
    r3.font.name = "Calibri"
    r3.font.color.rgb = RGBColor(30, 58, 138)
    
    p = doc.add_paragraph(
        "When real users begin piloting campus software, communication must be institutionalized. "
        "Never run customer support on personal WhatsApp or unmonitored inboxes. These tools professionalize user onboarding:"
    )
    p.paragraph_format.space_after = Pt(8)
    
    t3 = doc.add_table(rows=5, cols=4)
    h3_cols = ["Platform", "Program Benefit", "Functional Utility", "Onboarding Method"]
    for j, h in enumerate(h3_cols):
        t3.cell(0, j).text = h
        
    support_data = [
        ("Freshworks for Startups\n(Freshdesk & Freshchat)", "$10,000 credits across all Freshworks products (100% free Yr 1)", "Multi-channel ticketing, in-app chat widget, knowledge base, WhatsApp business integration.", "Requires College Incubation Cell to sign up as an official Partner (Free, online setup)."),
        ("HubSpot for Startups", "75% to 90% discount on CRM & Service Hub", "Complete customer lifecycle: lead pipeline, sales tracking, automated email outreach, customer tickets.", "Apply via approved incubator affiliate code or early-stage student founder application."),
        ("Intercom for Startups", "95% discount for 12 months (~$65/mo vs $1,000+)", "Industry-standard in-app live chat bubble, Fin AI customer support agent, interactive onboarding tours.", "Direct application via Intercom Early Stage Program page (requires active website)."),
        ("Crisp for Startups", "Free tier + heavy startup discounts", "Lightweight, easy-to-embed live chat and shared team inbox for Next.js and web platforms.", "Apply on crisp.chat/startups with proof of early-stage incubation.")
    ]
    for i, row_data in enumerate(support_data):
        for j, text in enumerate(row_data):
            t3.cell(i+1, j).text = text
            
    format_table(t3, col_widths=[Inches(1.6), Inches(1.5), Inches(2.0), Inches(1.6)])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # -------------------------------------------------------------
    # SECTION 4
    # -------------------------------------------------------------
    h4 = doc.add_heading(level=1)
    r4 = h4.add_run("4. Indian Legal, Billing, Payments & Operations Stack")
    r4.font.name = "Calibri"
    r4.font.color.rgb = RGBColor(30, 58, 138)
    
    p = doc.add_paragraph(
        "Operating in India requires GST-compliant invoicing, payment gateways (UPI, cards, NetBanking), and corporate banking. "
        "Two flagship domestic programs exist specifically for student and pre-seed founders:"
    )
    p.paragraph_format.space_after = Pt(8)
    
    t4 = doc.add_table(rows=4, cols=4)
    h4_cols = ["Program", "Financial & SaaS Value", "Operational Utility", "Target Audience"]
    for j, h in enumerate(h4_cols):
        t4.cell(0, j).text = h
        
    ops_data = [
        ("Razorpay Rize\n(India's Premier Founder Club)", "Free payment gateway + zero fee processing on first ₹5-10 Lakh; corporate card access.", "Collecting pilot subscription fees, automated vendor payouts, stipend disbursements.", "Specifically designed for early-stage and student founders in India."),
        ("Zoho for Startups", "₹1,86,000 (~$2,500) in Zoho Wallet credits", "Zoho Books (GST invoicing), Zoho Desk, Zoho CRM, custom domain email (Zoho Mail).", "Available when College Incubation Cell signs MoU with Zoho for Startups."),
        ("Stripe Atlas & Processing", "$20,000 fee-free global processing; incorporation assistance", "International SaaS subscription collection in USD/EUR from global contractors or clients.", "Direct application with accelerator/incubator code.")
    ]
    for i, row_data in enumerate(ops_data):
        for j, text in enumerate(row_data):
            t4.cell(i+1, j).text = text
            
    format_table(t4, col_widths=[Inches(1.6), Inches(1.6), Inches(2.0), Inches(1.5)])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # -------------------------------------------------------------
    # SECTION 5
    # -------------------------------------------------------------
    h5 = doc.add_heading(level=1)
    r5 = h5.add_run("5. Government & University Grants: Non-Dilutive Seed Capital")
    r5.font.name = "Calibri"
    r5.font.color.rgb = RGBColor(30, 58, 138)
    
    p = doc.add_paragraph(
        "In addition to software credits, campus startups under Savitribai Phule Pune University (SPPU) and "
        "Maharashtra state have access to government grant-in-aid schemes that provide non-dilutive prototype capital:"
    )
    p.paragraph_format.space_after = Pt(8)
    
    t5 = doc.add_table(rows=5, cols=4)
    h5_cols = ["Scheme / Grant", "Funding Ceiling", "Eligible Expenditures", "Nodal Channel in Pune / SPPU"]
    for j, h in enumerate(h5_cols):
        t5.cell(0, j).text = h
        
    gov_data = [
        ("DST NIDHI-PRAYAS\n(Dept. of Science & Tech)", "Up to ₹10,00,000 (100% Grant)", "Hardware/software prototyping, raw materials, testing, third-party fabrication.", "COEP Bhau Institute, NCL Venture Center, or partner college TBIs in Pune."),
        ("MeitY TIDE 2.0\n(ICT & Deep Tech)", "₹4,00,000 (Grant) to ₹7,00,000 (Scale-up)", "Software development, AI model training, cloud compute, testing with pilot users.", "MeitY approved incubation centers (COEP, PICT, etc.)."),
        ("MSInS Seed Grant\n(Govt of Maharashtra)", "Up to ₹5,00,000 per startup", "Proof of Concept validation, pilot trial execution, intellectual property filing.", "Maharashtra State Innovation Society direct portal."),
        ("Patent Support Scheme\n(MSInS / DPIIT)", "Up to ₹2,00,000 (Domestic) / ₹10,00,000 (Intl)", "Patent attorney fees, prior-art search, filing & examination fee reimbursement.", "DPIIT recognized startups via Startup India portal.")
    ]
    for i, row_data in enumerate(gov_data):
        for j, text in enumerate(row_data):
            t5.cell(i+1, j).text = text
            
    format_table(t5, col_widths=[Inches(1.6), Inches(1.4), Inches(2.0), Inches(1.7)])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # -------------------------------------------------------------
    # SECTION 6
    # -------------------------------------------------------------
    h6 = doc.add_heading(level=1)
    r6 = h6.add_run("6. Institutional Onboarding Roadmap for the E-Cell President")
    r6.font.name = "Calibri"
    r6.font.color.rgb = RGBColor(30, 58, 138)
    
    p = doc.add_paragraph(
        "To transform your college E-Cell into a recognized startup hub, execute this 4-step institutional partnership roadmap:"
    )
    p.paragraph_format.space_after = Pt(6)
    
    steps = [
        ("Phase 1: Institutional Partner Registration (Zero Cost)", 
         "Submit applications for your E-Cell to become an official partner for: (a) AWS Activate Provider, (b) Freshworks for Startups, and (c) Zoho for Startups. Each program only asks for college registration documents, incubation website/page link, and coordinator contact details."),
        ("Phase 2: Establish the Org ID Distribution Mechanism", 
         "Once approved, AWS and Freshworks will issue an exclusive 'Partner Org ID / Referral Code' to your college. Create an internal E-Cell form where vetted student startups submit their pitch deck and prototype link to receive the Org ID."),
        ("Phase 3: Host a 'Zero-Dollar Cloud Runway' Workshop", 
         "Conduct a hands-on session guiding student teams through Microsoft Founders Hub ($150k pipeline) and MongoDB for Startups. This produces immediate technical milestones for projects across AI, Web3, and Full-Stack domains."),
        ("Phase 4: Connect with Regional Pune TBIs for Grants", 
         "Establish academic outreach with COEP Bhau Institute and Venture Center Pune so eligible final-year projects (like OptiBuild) can fast-track applications for DST NIDHI-PRAYAS (₹10 Lakh) and MeitY TIDE 2.0 grants.")
    ]
    
    for title, desc in steps:
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(3)
        sp.paragraph_format.space_after = Pt(4)
        r_step = sp.add_run(f"🔹 {title}\n")
        r_step.bold = True
        r_step.font.name = "Calibri"
        r_step.font.color.rgb = RGBColor(37, 99, 235)
        r_desc = sp.add_run(desc)
        r_desc.font.name = "Calibri"
        r_desc.font.size = Pt(9.5)
        
    # -------------------------------------------------------------
    # SECTION 7
    # -------------------------------------------------------------
    h7 = doc.add_heading(level=1)
    r7 = h7.add_run("7. Ready-to-Send Outreach Templates")
    r7.font.name = "Calibri"
    r7.font.color.rgb = RGBColor(30, 58, 138)
    
    p = doc.add_paragraph(
        "Copy and adapt these communication templates for reaching out to college administration and startup program managers:"
    )
    p.paragraph_format.space_after = Pt(6)
    
    # Template Box 1: College Administration
    add_callout(
        doc,
        "Subject: Proposal to Register Our College Incubation Cell for Global Startup Credit Partnerships (AWS, Freshworks, Zoho)\n\n"
        "Respected Principal / Dean of R&D / Incubation Director,\n\n"
        "As President of the Entrepreneurship Cell (E-Cell), I am writing to propose registering our college's Incubation Centre as an official institutional partner with leading startup ecosystems, including AWS Activate, Freshworks for Startups, and Zoho for Startups.\n\n"
        "Key Institutional Benefits at ZERO Cost to the College:\n"
        "1. Up to $25,000 in AWS Cloud compute credits per student venture (vs $1,000 for unverified teams).\n"
        "2. $10,000 in Freshworks customer support and helpdesk tools per startup.\n"
        "3. ₹1,86,000 in Zoho credits for GST billing, enterprise email, and CRM.\n"
        "4. Tangible NIRF, NAAC, and NBA metrics under Student Innovation & Incubation.\n\n"
        "These programs require no financial investment from the institution—only an official digital application with college accreditation details. With your permission, the E-Cell team can prepare and coordinate the documentation.\n\n"
        "Looking forward to your guidance.\n\n"
        "Sincerely,\n"
        "[Your Name]\n"
        "President, Entrepreneurship Cell (E-Cell)\n"
        "[College Name, Pune]",
        "TEMPLATE A: PROPOSAL TO COLLEGE PRINCIPAL / INCUBATION DIRECTOR"
    )
    
    # Template Box 2: AWS Activate / Freshworks Program Manager
    add_callout(
        doc,
        "Subject: Institutional Partnership Application - [College Name] E-Cell & Incubation Center\n\n"
        "Dear Program Manager,\n\n"
        "I am writing on behalf of the Entrepreneurship Cell and Incubation Centre of [College Name, Pune, affiliated with SPPU]. Our institution supports over 50+ student innovators and active technology ventures in AI, Construction-Tech, Healthcare, and SaaS.\n\n"
        "We wish to onboard our incubation facility as an official Partner for your Startup Program to empower our student ventures with cloud infrastructure and tooling credits.\n\n"
        "Incubation Details:\n"
        "- Institution: [College Name]\n"
        "- Location: Pune, Maharashtra, India\n"
        "- Active Incubated Ventures: [e.g., 8-15 active student teams]\n"
        "- Official Website: [College/E-Cell URL]\n\n"
        "Please let us know the requisite verification steps or MoU documentation to formalize our institutional partner code.\n\n"
        "Warm regards,\n"
        "[Your Name]\n"
        "President, E-Cell & Student Incubation Lead",
        "TEMPLATE B: INQUIRY TO GLOBAL STARTUP PROGRAM MANAGERS"
    )
    
    # Final footer/signoff note
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("--- End of Master Playbook | Generated for SPPU E-Cell Leadership ---")
    r_end.font.name = "Calibri"
    r_end.font.size = Pt(9)
    r_end.font.italic = True
    r_end.font.color.rgb = RGBColor(148, 163, 184)
    
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    out_dir = r"d:\Opti Build\docs"
    os.makedirs(out_dir, exist_ok=True)
    out_file = os.path.join(out_dir, "ECell_Startup_Credits_Master_Playbook.docx")
    build_document(out_file)
